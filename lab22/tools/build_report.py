"""Сборка отчета DOCX. Нужна python-docx; программа и тесты от нее не зависят."""

from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output"


def math(paragraph, text):
    """Редактируемая линейная формула Word без растровых изображений."""
    equation = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    props = OxmlElement("m:rPr")
    style = OxmlElement("m:sty")
    style.set(qn("m:val"), "p")
    props.append(style)
    run.append(props)
    value = OxmlElement("m:t")
    value.text = text
    run.append(value)
    equation.append(run)
    paragraph._p.append(equation)


def table(doc, headings, rows, widths):
    tab = doc.add_table(rows=1, cols=len(headings))
    tab.autofit = False
    for column, width in zip(tab.columns, widths):
        column.width = Cm(width)
    for row_index, values in enumerate([headings, *rows]):
        cells = tab.rows[0].cells if row_index == 0 else tab.add_row().cells
        for cell, value, width in zip(cells, values, widths):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = str(value)
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                item = OxmlElement("w:" + side)
                item.set(qn("w:w"), "85")
                item.set(qn("w:type"), "dxa")
                margins.append(item)
            tcpr.append(margins)
            borders = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                edge = OxmlElement("w:" + side)
                edge.set(qn("w:val"), "single")
                edge.set(qn("w:sz"), "4")
                edge.set(qn("w:color"), "D9D9D9")
                borders.append(edge)
            tcpr.append(borders)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "E7EBEF" if row_index == 0 else ("F5F5F5" if row_index % 2 == 0 else "FFFFFF"))
            tcpr.append(shading)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.size = Pt(10.5)
                    run.bold = row_index == 0
        pr = tab.rows[row_index]._tr.get_or_add_trPr()
        prevent = OxmlElement("w:cantSplit")
        pr.append(prevent)
        if row_index == 0:
            repeat = OxmlElement("w:tblHeader")
            pr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return tab


def paragraph(doc, text):
    return doc.add_paragraph(text)


def formula(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    math(p, text)
    return p


def code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.0
    for index, line in enumerate(text.splitlines()):
        if index:
            p.add_run().add_break()
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(10)
    return p


def page(doc, title):
    doc.add_page_break()
    doc.add_heading(title, 1)


def main():
    test_log = (ROOT / "evidence/tests_python312.txt").read_text(encoding="utf-8")
    count = int(re.search(r"Ran (\d+) tests", test_log).group(1))
    if not test_log.rstrip().endswith("OK"):
        raise RuntimeError("Нет успешного протокола тестов")
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    sec.left_margin, sec.right_margin = Cm(2.2), Cm(1.8)
    sec.footer_distance = Cm(0.8)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(12)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.15
        # Шаблон python-docx содержит тематические шрифты и синюю линию Title.
        # Убираем наследованные украшения в исходном генераторе.
        for fonts in style.element.iter(qn("w:rFonts")):
            for key in list(fonts.attrib):
                if "theme" in key.lower():
                    del fonts.attrib[key]
            for key in ("ascii", "hAnsi", "eastAsia", "cs"):
                fonts.set(qn("w:" + key), "Times New Roman")
        for borders in list(style.element.iter(qn("w:pBdr"))):
            borders.getparent().remove(borders)
    doc.styles["Title"].font.size = Pt(22)
    doc.styles["Title"].font.bold = True
    doc.styles["Heading 1"].font.size = Pt(16)
    doc.styles["Heading 2"].font.size = Pt(13)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(4)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Страница ").font.size = Pt(10)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = "Лабораторная работа 1 Верификатор доказательств Вариант 3"
    doc.core_properties.subject = "Дискретная математика Неклассические логические исчисления"
    doc.core_properties.author = ""

    p = paragraph(doc, "Национальный исследовательский ядерный университет «МИФИ»")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = paragraph(doc, "Дискретная математика\nНеклассические логические исчисления")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    p = doc.add_paragraph("Лабораторная работа 1", "Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Верификатор доказательств теорем\nВариант 3", "Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    p = paragraph(doc, "Москва 2026")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("Цель работы", 1)
    paragraph(doc, "Разработать консольную программу, которая проверяет каждую вводимую строку доказательства в системе исчисления высказываний. Формула принимается, если она является исходной аксиомой, уже доказана либо получается из доступных формул одним применением modus ponens или одной равномерной подстановкой β.")
    doc.add_heading("Результат", 1)
    paragraph(doc, "Реализована программа на Python 3.10+ без внешних зависимостей для запуска и тестирования. Предусмотрены интерактивный и пакетный режимы, печать дерева, список доказанных строк и очистка сеанса. Для каждого принятого шага сохраняется объяснение с источниками.")
    paragraph(doc, f"Проверка на Python 3.12.14 и 3.14.3 завершилась успешно: {count} автоматических теста на каждой версии. Демонстрационное доказательство содержит семь принятых строк. Отрицательный набор содержит восемь отклоненных формул.")
    paragraph(doc, "До реализации изучены 26 исходных файлов: 58 слайдов, пять страниц PDF и 21 изображение. Подробный разбор и перечень контрольных сумм включены в комплект работы.")

    page(doc, "1 Формальная система варианта 3")
    paragraph(doc, "Базовые связки — отрицание и дизъюнкция. Переменная обозначается одной латинской буквой. Формула строится из переменной, отрицания формулы или дизъюнкции двух формул. Буквы f и t в этом варианте являются обычными переменными: специальные константы не заданы.")
    formula(doc, "A ⊃ B := ¬A ∨ B")
    paragraph(doc, "Импликация используется только как сокращение. Ее раскрытие выполняется при разборе; семантические преобразования дерева не выполняются. Исходные аксиомы заданы конкретными формулами:")
    axioms = (
        "A1   p ⊃ (p ∨ q)",
        "A2   (p ∨ p) ⊃ p",
        "A3   (p ∨ (q ∨ r)) ⊃ (q ∨ (p ∨ r))",
        "A4   (q ⊃ s) ⊃ ((p ∨ q) ⊃ (p ∨ s))",
        "A5   ¬¬p ⊃ p",
        "A6   (¬q ⊃ ¬p) ⊃ (p ⊃ q)",
        "A7   ¬p ∨ p",
    )
    for text in axioms:
        formula(doc, text)
    paragraph(doc, "В пункте 7 задания после формулы напечатано равенство с другой записью через импликации. В набор включена левая формула ¬p∨p. Знак равенства не входит в язык ввода, правая запись не добавляется как отдельная аксиома. При принятом раскрытии импликации это разные деревья.")
    doc.add_heading("Правила вывода", 2)
    formula(doc, "MP   A, A ⊃ B ⊢ B")
    formula(doc, "β   A ⊢ [X/x]A")
    paragraph(doc, "В β заменяются все вхождения ровно одной переменной x на одно дерево X. Вставляемая формула может быть составной. Одновременная замена разных переменных требует нескольких строк доказательства.")
    paragraph(doc, "Например, p⊃p раскрывается в ¬p∨p и совпадает с A7. Длинный пример K/S из методички относится к другому варианту и не определяет набор аксиом этой реализации.")

    page(doc, "2 Синтаксический анализ и дерево")
    paragraph(doc, "Лексер сохраняет позиции токенов в исходной строке и распознает однобуквенные переменные, операции и скобки. Математические варианты латинских букв из PDF нормализуются по NFKC посимвольно. Пробелы игнорируются, но соседние буквы не объединяются в имя переменной.")
    table(doc, ["Операция", "Запись"], [
        ["Отрицание", "¬, ~, !"], ["Дизъюнкция", "∨, |"], ["Импликация", "⊃, →, ->"],
        ["Группировка", "Круглые скобки; лишние парные скобки допустимы"],
    ], [4.3, 12.7])
    paragraph(doc, "Разбор реализован рекурсивным спуском. Грамматика EBNF использует повторение в фигурных скобках и выбор через вертикальную черту:")
    code(doc, 'formula     = disjunction { IMP disjunction }\ndisjunction = unary { OR unary }\nunary       = VARIABLE | NOT unary | "(" formula ")"')
    paragraph(doc, "Приоритет убывает от отрицания к дизъюнкции и затем к импликации. Бинарные связки группируются слева. Для импликации это соответствует лекции 1, слайд 16:")
    formula(doc, "a ⊃ b ⊃ c = (a ⊃ b) ⊃ c ≠ a ⊃ (b ⊃ c)")
    paragraph(doc, "Неизменяемое дерево использует три типа узлов: Var(name), Not(operand), Or(left, right). Узла скобок нет. Например, команда :tree p -> q печатает:")
    code(doc, "OR\n  NOT\n    p\n  q")
    paragraph(doc, "Структурное равенство требует одинаковых операций и одинаковых соответствующих поддеревьев. Перестановка дизъюнктов, иная группировка и удаление двойного отрицания не выполняются. Поэтому логически эквивалентные выражения не обязательно совпадают в верификаторе.")
    paragraph(doc, "Ошибки содержат позицию: недопустимый символ, отсутствующий операнд, скобка или операция. Для защиты от чрезмерного ввода установлены пределы 4096 символов, 511 узлов после раскрытия и глубина 80.")

    page(doc, "3 Проверка шага доказательства")
    paragraph(doc, "Сеанс хранит семь аксиом, принятые строки Step и их основания Reason. Источник Reference указывает номер аксиомы или более ранней строки. Доступные деревья индексируются словарем; для MP дополнительно строится индекс по консеквенту. Ошибочная строка не изменяет ни доказательство, ни индексы.")
    doc.add_heading("Порядок проверки", 2)
    for text in (
        "1. Разобрать ввод и проверить ограничения. При синтаксической ошибке вернуть отказ.",
        "2. Проверить точное совпадение с аксиомой или ранее доказанной формулой. Для повтора сохранить ссылку на первое основание.",
        "3. Найти в индексе импликацию с консеквентом B и проверить наличие ее антецедента A среди доступных деревьев.",
        "4. Для каждого доступного источника и каждой его переменной проверить, является ли B одной равномерной подстановкой.",
        "5. Если основание найдено, добавить строку с объяснением. Иначе сообщить, что непосредственный шаг не обоснован.",
    ):
        paragraph(doc, text)
    doc.add_heading("Сопоставление подстановки", 2)
    paragraph(doc, "Для выбранной переменной x параллельно обходятся источник и результат. При первом вхождении x запоминается соответствующее поддерево X результата. На следующих вхождениях требуется точно такое же X. Все остальные переменные и связки обязаны совпасть. Несогласованность отвергает кандидата. Это позволяет найти вставку без перебора всех возможных формул.")
    formula(doc, "[p ∨ q / p](p ∨ p) = (p ∨ q) ∨ (p ∨ q)")
    paragraph(doc, "Вставленное дерево повторно не обходится. Допускается наличие x в X. Частичная замена x и изменение другой переменной в том же шаге недопустимы. Простая замена подстроки вместо обхода дерева не используется.")
    doc.add_heading("Корректность проверки", 2)
    paragraph(doc, "Индукция по принятым строкам: исходные аксиомы разрешены по определению; повтор ссылается на уже установленный факт; MP проверяет обе посылки; сопоставление β строит свидетельство одной равномерной подстановки. Поэтому каждая новая строка имеет разрешенное основание. Будущие и отклоненные строки недоступны.")
    paragraph(doc, "Алгоритм исчерпывает основания одного шага относительно текущего конечного набора. Он не ищет произвольное многошаговое доказательство. Отказ не утверждает недоказуемость формулы во всей системе. Семантическая эквивалентность и таблицы истинности не используются для приема строк.")

    page(doc, "4 Демонстрационное доказательство")
    paragraph(doc, "Файл examples/demo.txt содержит последовательность из семи формул. В ней отдельно показаны β из аксиомы, β из предыдущей строки, составная вставка и два применения MP. Аксиомы доступны без предварительного ручного ввода.")
    entries = [
        ("¬q ⊃ ¬q", "β из A7: p заменяется на ¬q."),
        ("(¬q ⊃ ¬q) ⊃ (q ⊃ q)", "β из A6: p заменяется на q."),
        ("q ⊃ q", "MP из строк 2 и 1."),
        ("p ⊃ (p ∨ r)", "β из A1: q заменяется на r."),
        ("(q ⊃ q) ⊃ ((q ⊃ q) ∨ r)", "β из строки 4: p заменяется на (q ⊃ q)."),
        ("(q ⊃ q) ∨ r", "MP из строк 5 и 3."),
        ("(q ⊃ q) ∨ (s ∨ ¬s)", "β из строки 6: r заменяется на (s ∨ ¬s)."),
    ]
    for index, (expression, reason) in enumerate(entries, 1):
        formula(doc, f"{index}.   {expression}")
        p = paragraph(doc, reason)
        p.paragraph_format.space_after = Pt(11)
    paragraph(doc, "В строке 3 существует и другое основание — β из A7. В текущем сеансе программа сначала обнаруживает MP. Наличие нескольких корректных оснований не влияет на результат проверки.")
    paragraph(doc, "Реальный итог запуска: «Проверено: 7. Принято: 7. Отклонено: 0». Протоколы сохранены в evidence/demo_python312.txt и evidence/demo_python314.txt. Формулы результата печатаются после раскрытия импликаций в базисе ¬ и ∨.")

    page(doc, "5 Тестирование")
    paragraph(doc, f"Выполнены {count} теста стандартного unittest на Windows 11 с Python 3.12.14 и Python 3.14.3. Дополнительно проверены компиляция исходников, встроенная демонстрация, ее запуск из файла и отрицательный файл. Команда tools/verify.py сохраняет результаты и коды выхода.")
    table(doc, ["Проверка", "Фактический результат"], [
        ["Полный набор unittest", f"{count}/{count} на каждой из двух версий Python"],
        ["Семь исходных аксиом", "Все приняты; все общезначимы в независимом булевом тесте"],
        ["Демонстрация и тот же файл", "7 из 7, код выхода 0"],
        ["Отрицательный файл", "0 из 8, код выхода 1 как ожидается"],
        ["Ошибки файла и UTF-8", "Диагностика и код выхода 2"],
        ["Подстановка", "Сопоставление проверено также на 675 конечных комбинациях"],
        ["Состояние и ссылки", "Отказы не добавляются, сеансы независимы, ссылки ведут назад"],
        ["Исходные материалы", "SHA-256 всех 26 файлов совпали с первоначальным перечнем"],
    ], [6.0, 11.0])
    doc.add_heading("Проверки отказов", 2)
    paragraph(doc, "Из методички проверены лишняя скобка, соседние переменные mn, пропущенная операция между p p и импликация без левого или правого операнда. Корректный пример с лишними парными скобками успешно разбирается. Проверены математические Unicode-буквы, приоритеты и левая группировка.")
    paragraph(doc, "Логические отрицательные случаи: две замены вместо одной, замена только части вхождений, отсутствие антецедента MP, попытка получить посылку из заключения, автоматическая перестановка дизъюнктов и добавление неверной строки в доступную базу. В примере a⊃(a∨b) первый шаг отвергается, но после a⊃(a∨q) результат принимается отдельной заменой q на b.")
    doc.add_heading("Границы проверки", 2)
    paragraph(doc, "Независимые булевы оценки служат только дополнительной проверкой аксиом и демонстрации. Они не являются математическим доказательством отсутствия всех программных ошибок. Запуск на другом физическом компьютере, на Linux/macOS и на Python 3.10/3.11/3.13 отдельно не выполнялся. Для проекта не требуется сеть или установленный офисный пакет.")

    page(doc, "6 Запуск и выводы")
    paragraph(doc, "Для работы достаточно распаковать комплект и открыть терминал в его корневой папке. В Windows используется команда py -3; при отсутствии Python Launcher — python. Внешние пакеты устанавливать не нужно.")
    code(doc, "py -3 lab1.py\npy -3 lab1.py --demo\npy -3 lab1.py --file examples/demo.txt\npy -3 -m unittest discover -s tests -v\npy -3 tools/verify.py")
    paragraph(doc, "Интерактивные команды: :help, :axioms, :list, :tree ФОРМУЛА, :reset, :quit. Пустые строки и комментарии с # пропускаются. В пакетном режиме вводится одна формула на строку; файл читается как UTF-8. Максимум сеанса — 1000 принятых строк, после чего можно начать новый сеанс.")
    table(doc, ["Файл или каталог", "Назначение"], [
        ["lab1.py, proofchecker/", "Точка входа, дерево, парсер, аксиомы, правила и сеанс"],
        ["tests/", "Тесты синтаксиса, подстановки, MP, состояния и CLI"],
        ["examples/", "Готовые положительный и отрицательный наборы"],
        ["docs/", "Разбор материалов, соответствие требованиям, сценарий защиты"],
        ["evidence/", "Фактические протоколы и машиночитаемые итоги"],
    ], [6.0, 11.0])
    doc.add_heading("Вывод", 2)
    paragraph(doc, "Задача выполнена: создан верификатор доказательств для варианта 3 с абстрактным синтаксическим деревом, структурным сравнением, MP и атомарным β. Обоснованные формулы сохраняются с источниками, ошибочный ввод и необоснованные шаги отвергаются. Логика программы отделена от консольного интерфейса и проверяется автоматическими тестами.")
    doc.add_heading("Использованные материалы", 2)
    for text in (
        "1. Лабораторная работа №1.pdf, с. 1–3, требования и вариант 3.",
        "2. Методический материал к ЛР.pdf, с. 1–2, примеры синтаксиса и сеанса.",
        "3. Лекция 1.pptx, 25 слайдов, формулы, доказательства, MP, подстановка.",
        "4. Лекция 2-1.pptx, 20 слайдов, выводимость, семантика, многозначные логики.",
        "5. Семинар №1 2024.pptx, 13 слайдов; пять PNG и 16 фотографий доски.",
    ):
        p = paragraph(doc, text)
        p.paragraph_format.space_after = Pt(4)
        for run in p.runs:
            run.font.size = Pt(10.5)
    OUT.mkdir(exist_ok=True)
    path = OUT / "LR1_variant3.docx"
    doc.save(path)
    print(path)


if __name__ == "__main__":
    main()
